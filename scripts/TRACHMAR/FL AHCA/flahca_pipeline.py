"""FL AHCA preparation, Bulk Mailer reconciliation, and archive stages."""

import argparse
from contextlib import contextmanager
import csv
from datetime import date, datetime
import hashlib
import io
import json
import os
from pathlib import Path
import re
import secrets
import tempfile
import time
import zipfile

BASE = Path(r"C:\Goji\AUTOMATION\TRACHMAR\FL AHCA")
FIELD_MAP = {
    "MEM_First_Name": "First Name",
    "MEM_Last_Name": "Last Name",
    "MEM_Addr_1": "Address Line 1",
    "MEM_Addr_2": "Address Line 2",
    "MEM_City": "City",
    "MEM_State": "State",
    "MEM_Zip": "ZIP Code",
}
ID = "UNIQUE ID"
INPUT_HEADERS = [ID, *FIELD_MAP.values()]
MAPPING = "TMFLAHCA SOURCE.csv"
MANIFEST = "TMFLAHCA JOB.json"
EXPORT = "TRACHMAR FL AHCA.csv"
CODE_LIST = "TMFLAHCA 14 CODE LIST.csv"


class JobError(Exception):
    pass


def digest(content):
    return hashlib.sha256(content).hexdigest()


def csv_bytes(headers, rows):
    text = io.StringIO(newline="")
    writer = csv.writer(text)
    writer.writerow(headers)
    writer.writerows(rows)
    return text.getvalue().encode("utf-8-sig")


def json_bytes(value):
    return (json.dumps(value, indent=2, ensure_ascii=False) + "\n").encode("utf-8")


def read_csv(path):
    content = path.read_bytes()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    table = list(csv.reader(io.StringIO(text, newline="")))
    if not table or not table[0] or any(not h.strip() for h in table[0]):
        raise JobError(f"Missing column headers: {path.name}")
    headers = table[0]
    if len(set(headers)) != len(headers):
        raise JobError(f"Duplicate column headers: {path.name}")
    rows = [r for r in table[1:] if any(v.strip() for v in r)]
    if any(len(r) != len(headers) for r in rows):
        raise JobError(f"Incorrect column count in {path.name}")
    return headers, rows


def publish(outputs, *, replace=False):
    """Stage all files first; restore existing results if publication fails."""
    staged = {}
    previous = {}
    committed = []
    try:
        for path, content in outputs.items():
            path.parent.mkdir(parents=True, exist_ok=True)
            previous[path] = path.read_bytes() if path.exists() else None
            if previous[path] is not None and previous[path] != content and not replace:
                raise JobError(f"Refusing to replace existing file: {path}")
            with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".flahca-", delete=False) as f:
                staged[path] = Path(f.name)
                f.write(content)
                f.flush()
                os.fsync(f.fileno())
            if staged[path].read_bytes() != content:
                raise JobError(f"Staged file verification failed: {path.name}")
        for path, temp in staged.items():
            os.replace(temp, path)
            committed.append(path)
    except Exception:
        for path in reversed(committed):
            if previous[path] is None:
                path.unlink()
            else:
                path.write_bytes(previous[path])
        raise
    finally:
        for temp in staged.values():
            if temp.exists():
                temp.unlink()


@contextmanager
def job_lock(base):
    base.mkdir(parents=True, exist_ok=True)
    lock = base / ".flahca-processing.lock"
    try:
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise JobError("FL AHCA is already processing, or a previous interrupted run left "
                       f"a lock. Check that no script is running before removing {lock}.") from exc
    try:
        with os.fdopen(fd, "w") as f:
            f.write(str(os.getpid()))
        yield
    finally:
        lock.unlink()


def cell_text(cell):
    value = cell.value
    if cell.data_type == "f":
        raise JobError("Source formulas are not supported; provide a workbook containing values.")
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_source(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    candidates = []
    try:
        for sheet in workbook.worksheets:
            first = next(sheet.iter_rows(), ())
            headers = [cell_text(c) for c in first]
            if set(FIELD_MAP).issubset(headers):
                candidates.append((sheet, headers))
        if len(candidates) != 1:
            raise JobError("Expected exactly one worksheet with all seven MEM_ address columns.")
        sheet, headers = candidates[0]
        if any(not h.strip() for h in headers) or len(headers) != len(set(headers)):
            raise JobError("Source contains blank or duplicate column headers.")
        if ID in headers or "MAILED" in headers:
            raise JobError("Source already contains a reserved UNIQUE ID or MAILED column.")
        rows = []
        for cells in sheet.iter_rows(min_row=2, max_col=len(headers)):
            values = [cell_text(c) for c in cells]
            if any(v != "" for v in values):
                rows.append(values)
        if not rows:
            raise JobError("Source workbook contains no records.")
        return sheet.title, headers, rows
    finally:
        workbook.close()


def input_rows(headers, rows, ids):
    positions = [headers.index(h) for h in FIELD_MAP]
    result = []
    for uid, source in zip(ids, rows):
        values = [source[i] for i in positions]
        # Excel often stores ZIPs numerically; write the mailing field as text.
        zipcode = values[-1]
        if zipcode.isascii() and zipcode.isdigit() and len(zipcode) <= 5:
            values[-1] = zipcode.zfill(5)
        result.append([uid, *values])
    return result


def load_job(base, identity):
    data = base / "DATA"
    manifest = json.loads((data / MANIFEST).read_text(encoding="utf-8"))
    if manifest.get("identity") != identity:
        raise JobError("Prepared data belongs to a different job, year, or month.")
    name = manifest.get("source_name", "")
    if not name or Path(name).name != name:
        raise JobError("Invalid source filename in job manifest.")
    source = base / "RAW INPUT" / name
    if digest(source.read_bytes()) != manifest["source_sha256"]:
        raise JobError("Source workbook changed after preparation. Original IDs cannot be reused.")
    for filename, key in [(MAPPING, "mapping_sha256"), ("INPUT.csv", "input_sha256")]:
        if digest((data / filename).read_bytes()) != manifest[key]:
            raise JobError(f"Prepared {filename} changed after initial processing.")
    headers, rows = read_csv(data / MAPPING)
    if headers != [ID, *manifest["source_headers"]] or len(rows) != manifest["record_count"]:
        raise JobError("Source mapping does not match the prepared workbook structure.")
    ids = validate_ids(headers, rows, MAPPING)
    expected_input = csv_bytes(INPUT_HEADERS, input_rows(headers[1:], [r[1:] for r in rows], ids))
    if (data / "INPUT.csv").read_bytes() != expected_input:
        raise JobError("INPUT.csv no longer corresponds to the original source mapping.")
    return manifest, source, headers, rows, ids


def validate_ids(headers, rows, label, allowed=None):
    if ID not in headers:
        raise JobError(f"{label} must contain a {ID} column exported unchanged from Bulk Mailer.")
    pos = headers.index(ID)
    ids = [r[pos].strip() for r in rows]
    if any(not re.fullmatch(r"[A-Z0-9]{10}", uid) for uid in ids):
        raise JobError(f"Blank or malformed UNIQUE ID in {label}.")
    if len(ids) != len(set(ids)):
        raise JobError(f"Duplicate UNIQUE ID in {label}.")
    if allowed is not None and not set(ids).issubset(allowed):
        raise JobError(f"Unknown UNIQUE ID in {label}; check that the export belongs to this job.")
    return ids


def prepare(base, identity):
    raw = base / "RAW INPUT"
    data = base / "DATA"
    files = sorted(p for p in raw.glob("*.xlsx") if not p.name.startswith("~$"))
    if len(files) != 1:
        raise JobError("Place exactly one source XLSX workbook in RAW INPUT.")
    if (data / MANIFEST).exists():
        manifest, source, _, _, _ = load_job(base, identity)
        if source != files[0]:
            raise JobError("RAW INPUT does not match the prepared source.")
        print(f"Existing INPUT.csv verified; preserved {manifest['record_count']} row IDs.")
        return
    if data.exists() and any(data.iterdir()):
        raise JobError("DATA contains files from an unprepared or earlier job. Preserve or remove "
                       "those files before running initial processing; nothing was overwritten.")
    source = files[0]
    source_content = source.read_bytes()
    sheet, headers, rows = read_source(source)
    ids = []
    used = set()
    while len(ids) < len(rows):
        uid = secrets.token_hex(5).upper()
        if uid not in used:
            used.add(uid)
            ids.append(uid)
    mapping = csv_bytes([ID, *headers], [[uid, *r] for uid, r in zip(ids, rows)])
    mailing = csv_bytes(INPUT_HEADERS, input_rows(headers, rows, ids))
    manifest = {"version": 1, "identity": identity, "source_name": source.name,
                "sheet": sheet, "source_headers": headers, "record_count": len(rows),
                "source_sha256": digest(source_content), "mapping_sha256": digest(mapping),
                "input_sha256": digest(mailing)}
    if source.read_bytes() != source_content:
        raise JobError("Source workbook changed while it was being read.")
    publish({data / MAPPING: mapping, data / "INPUT.csv": mailing,
             data / MANIFEST: json_bytes(manifest)})
    print(f"Created INPUT.csv with {len(rows)} records and stable UNIQUE ID values.")
    print("Original workbook preserved in RAW INPUT.")


def reconcile(base, identity):
    manifest, source, headers, originals, ids = load_job(base, identity)
    data = base / "DATA"
    final_headers, final_rows = read_csv(data / EXPORT)
    code_headers, code_rows = read_csv(data / CODE_LIST)
    allowed = set(ids)
    mailed = set(validate_ids(final_headers, final_rows, EXPORT, allowed))
    not_mailed = set(validate_ids(code_headers, code_rows, CODE_LIST, allowed))
    if mailed & not_mailed:
        raise JobError("A UNIQUE ID appears in both the mailing export and the 14 code list.")
    if mailed | not_mailed != allowed:
        raise JobError("Bulk Mailer exports do not account for every source record.")
    status_columns = [i for i, header in enumerate(code_headers)
                      if header.strip().casefold() == "user text 3"]
    if len(status_columns) > 1:
        raise JobError("The 14 code list contains duplicate User text 3 columns.")
    if status_columns:
        pos = status_columns[0]
        if any(r[pos].strip() != "14" for r in code_rows):
            raise JobError("The 14 code list contains a User text 3 value other than 14.")
    if "Sort Position" not in final_headers:
        raise JobError(f"{EXPORT} is missing Sort Position.")
    pos = final_headers.index("Sort Position")
    positions = [r[pos].strip() for r in final_rows]
    if any(not re.fullmatch(r"[0-9]+", p) or int(p) < 1 for p in positions):
        raise JobError("Sort Position must contain positive whole numbers.")
    if len(set(map(int, positions))) != len(positions):
        raise JobError("Duplicate Sort Position in the mailing export.")
    merged = csv_bytes(headers[1:] + ["MAILED"],
                       [r[1:] + ["14" if r[0] in not_mailed else "13"] for r in originals])
    return manifest, source, merged, max(map(int, positions), default=0)


def find_print_folder(network_root, job):
    if not network_root.is_dir():
        raise OSError(f"Network job directory is unavailable: {network_root}")
    pattern = re.compile(rf"^{re.escape(job)}(?![0-9])")
    matches = [p for p in network_root.iterdir() if p.is_dir() and pattern.search(p.name)]
    if len(matches) > 1:
        raise JobError(f"More than one network folder matches job {job}; resolve the duplicate folders before continuing.")
    job_folder = matches[0] if matches else network_root / f"{job}_FLAHCA"
    for subfolder in ("PDF for Client", "Files for Ricoh", "Original Files",
                      "HP Indigo/DATA", "HP Indigo/PRINT", "HP Indigo/PROOF"):
        (job_folder / subfolder).mkdir(parents=True, exist_ok=True)
    return job_folder / "HP Indigo" / "DATA"


def deliver_print(content, identity, maximum, delivery_dir=None, network_root=None, fallback_dir=None):
    job, year, _ = identity
    filename = f"{job} TRACHMAR FL AHCA_{maximum}.csv"
    if delivery_dir is not None:
        destination = delivery_dir / filename
        publish({destination: content})
        return destination
    network_root = network_root if network_root is not None else Path(
        rf"\\NAS1069D9\AMPrintData\{year}_SrcFiles\T\Trachmar")
    intended = network_root / f"{job}_FLAHCA" / "HP Indigo" / "DATA"
    for attempt in range(3):
        try:
            intended = find_print_folder(network_root, job)
            destination = intended / filename
            publish({destination: content})
            if destination.read_bytes() != content:
                raise OSError("Print file verification failed after network copy.")
            print(f"Print file copied and verified on server: {destination}")
            return destination
        except OSError as exc:
            print(f"WARNING: Network copy attempt {attempt + 1}/3 failed: {exc}")
            if attempt < 2:
                time.sleep(2)
    fallback = fallback_dir if fallback_dir is not None else Path(r"C:\Users\JCox\Desktop\MOVE TO NETWORK DRIVE")
    destination = fallback / filename
    publish({destination: content,
             destination.with_suffix(".txt"): f"Intended network location: {intended}\n".encode("utf-8")})
    print(f"WARNING: Server delivery unavailable. Print file retained in fallback folder: {destination}")
    return destination


def prearchive(base, identity, delivery_dir=None, network_root=None, fallback_dir=None):
    data = base / "DATA"
    exports = {name: (data / name).read_bytes() for name in (EXPORT, CODE_LIST)}
    manifest, source, merged, maximum = reconcile(base, identity)
    if any((data / name).read_bytes() != content for name, content in exports.items()):
        raise JobError("Exports changed during validation; rerun Final Step.")
    merged_path = data / (source.stem + "_MERGED.csv")
    print_content = exports[EXPORT]
    destination = deliver_print(print_content, identity, maximum, delivery_dir, network_root, fallback_dir) if maximum else None
    if any((data / name).read_bytes() != content for name, content in exports.items()):
        raise JobError("Exports changed during print delivery; no email handoff or cleanup was performed.")
    manifest["prepared"] = {"export_sha256": digest(exports[EXPORT]),
                            "code_list_sha256": digest(exports[CODE_LIST]),
                            "merged_name": merged_path.name, "merged_sha256": digest(merged),
                            "max_sort_position": maximum,
                            "print_path": str(destination.resolve()) if destination else "",
                            "print_sha256": digest(print_content) if destination else ""}
    publish({merged_path: merged, data / MANIFEST: json_bytes(manifest)}, replace=True)
    print(f"Created and verified {merged_path.name}; {manifest['record_count']} source records retained.")
    print("=== NAS_FOLDER_PATH ===")
    print(str(data.resolve()))
    print("=== END_NAS_FOLDER_PATH ===")
    print("=== PRINT_FOLDER_PATH ===")
    print(str(destination.parent.resolve()) if destination else "")
    print("=== END_PRINT_FOLDER_PATH ===")
    print("=== PAUSE_FOR_EMAIL ===", flush=True)


def archive(base, identity):
    manifest, source, merged, maximum = reconcile(base, identity)
    data = base / "DATA"
    prepared = manifest.get("prepared")
    if not prepared:
        raise JobError("Run the prearchive phase and attach the merged CSV before archiving.")
    expected = {"export_sha256": digest((data / EXPORT).read_bytes()),
                "code_list_sha256": digest((data / CODE_LIST).read_bytes()),
                "merged_name": source.stem + "_MERGED.csv", "merged_sha256": digest(merged),
                "max_sort_position": maximum}
    if any(prepared.get(key) != value for key, value in expected.items()) or (data / prepared["merged_name"]).read_bytes() != merged:
        raise JobError("Exports or merged output changed after the email handoff; rerun Final Step.")
    if maximum:
        print_path = prepared.get("print_path")
        if not print_path:
            raise JobError("Run Final Step again to deliver the print file before archiving.")
        delivered = Path(print_path).read_bytes()
        if delivered != (data / EXPORT).read_bytes() or digest(delivered) != prepared.get("print_sha256"):
            raise JobError("Delivered print file changed after the email handoff; cleanup stopped.")
    if any(not p.is_file() for p in data.iterdir()):
        raise JobError("Unexpected subdirectory in DATA; archive stopped without cleanup.")
    payload = {"RAW INPUT/" + source.name: source.read_bytes()}
    payload.update({"DATA/" + p.name: p.read_bytes() for p in sorted(data.iterdir())})
    job, year, month = identity
    archive_path = base / "ARCHIVE" / f"{job} TM FL AHCA {month} {year} {manifest['source_sha256'][:10]}.zip"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in payload.items():
            z.writestr(name, content)
    if archive_path.exists():
        verify_archive(archive_path, payload)
    else:
        publish({archive_path: buffer.getvalue()})
    verify_archive(archive_path, payload)
    if not maximum:
        print("No mailed records; no print file was delivered.")
    paths = {name: base / name for name in payload}
    if any(p.read_bytes() != payload[name] for name, p in paths.items()):
        raise JobError("Job files changed during archiving; cleanup stopped.")
    try:
        for p in paths.values():
            p.unlink()
    except Exception:
        # Keep the validated ZIP and restore every missing input for an explicit retry.
        publish({p: payload[name] for name, p in paths.items() if not p.exists()})
        raise
    print(f"Verified archive created: {archive_path}")
    print("Completed job files cleared from DATA and RAW INPUT.")


def verify_archive(path, payload):
    with zipfile.ZipFile(path) as z:
        if set(z.namelist()) != set(payload) or z.testzip() is not None:
            raise JobError("Archive verification failed; original files were retained.")
        if any(z.read(name) != content for name, content in payload.items()):
            raise JobError("Archive contents do not match this job; original files were retained.")


def main(stage):
    parser = argparse.ArgumentParser()
    parser.add_argument("job_number")
    parser.add_argument("year")
    parser.add_argument("month")
    parser.add_argument("--base-dir", type=Path, default=BASE)
    if stage == "final":
        parser.add_argument("--mode", choices=["prearchive", "archive"], default="prearchive")
        parser.add_argument("--delivery-dir", type=Path)
        parser.add_argument("--network-root", type=Path)
        parser.add_argument("--fallback-dir", type=Path)
    args = parser.parse_args()
    if not re.fullmatch(r"[0-9]{5}", args.job_number):
        parser.error("Job number must contain five digits.")
    if not re.fullmatch(r"[0-9]{4}", args.year) or not args.month.isdigit() or not 1 <= int(args.month) <= 12:
        parser.error("Select a valid four-digit year and month 01 through 12.")
    identity = [args.job_number, args.year, f"{int(args.month):02d}"]
    try:
        with job_lock(args.base_dir):
            if stage == "initial":
                prepare(args.base_dir, identity)
            elif args.mode == "prearchive":
                prearchive(args.base_dir, identity, args.delivery_dir, args.network_root, args.fallback_dir)
            else:
                archive(args.base_dir, identity)
    except Exception as exc:
        print(f"ERROR: {exc}", flush=True)
        return 1
    return 0
