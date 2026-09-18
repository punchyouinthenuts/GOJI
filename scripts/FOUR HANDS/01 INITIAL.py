import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
import traceback
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd


print("=== INITIALIZING ===")

BASE_DIR = Path(r"C:\Goji\AUTOMATION\FOUR HANDS")
SOURCE_DIR = BASE_DIR / "ORIGINAL"
MANIFEST_FILE = BASE_DIR / ".goji_fourhands_state.json"

RES_INPUT_DIR = BASE_DIR / "RESIDENTIAL" / "INPUT"
HOSP_INPUT_DIR = BASE_DIR / "HOSPITALITY" / "INPUT"

RES_OUTPUT_FILE = RES_INPUT_DIR / "INPUT.csv"
HOSP_OUTPUT_FILE = HOSP_INPUT_DIR / "INPUT.csv"

DETECTION_MODE = "--detect-versions"
DETECTION_SOURCE_OPTION = "--source-dir"
DETECTION_BEGIN_MARKER = "=== FH_VERSION_DETECTION_BEGIN ==="
DETECTION_END_MARKER = "=== FH_VERSION_DETECTION_END ==="

SHEET_RESIDENTIAL = "Residential"
SHEET_COMMERCIAL = "Commercial"
SHEET_NEW_ADDRESSES = "New Addresses"

VERSION_RESIDENTIAL = "RESIDENTIAL"
VERSION_HOSPITALITY = "HOSPITALITY"
VERSION_ORDER = [VERSION_RESIDENTIAL, VERSION_HOSPITALITY]
SUPPORTED_SOURCE_EXTENSIONS = {".xlsx", ".xls", ".csv"}
SUPPORTED_RAW_EXTENSIONS = SUPPORTED_SOURCE_EXTENSIONS | {".zip"}
MAX_REPORTED_ZIP_WARNINGS = 10
MAX_HEADER_SCAN_ROWS = 500
MAX_TRAILING_BLANK_ROWS = 1000

IGNORE_MARKER_RE = re.compile(r"ignore\s*>\s*>\s*>", re.IGNORECASE)
STANDALONE_HOSPITALITY_RE = re.compile(
    r"^hospitality customer look books?\b",
    re.IGNORECASE,
)
RESIDENTIAL_SIGNAL_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:residential|brand)(?![A-Za-z0-9])",
    re.IGNORECASE,
)
HOSPITALITY_SIGNAL_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:hospitality|commercial)(?![A-Za-z0-9])",
    re.IGNORECASE,
)
COLUMNS_TO_REMOVE = [
    "Account Source",
    "Account Owner",
    "Contact Owner",
    "Account: Created Date↑",
    "Account: Created Date",
    "Customer Status",
    "Social Profile: Instagram",
]

COLUMN_RENAME = {
    "Account Name": "Business",
    "Billing Address Line 1": "Address Line 1",
    "Billing City": "City",
    "Billing State/Province": "State",
    "Billing Zip/Postal Code": "ZIP Code",
    "Mailing Country": "Country",
}

TEXT_ONLY_STATE_COL = "Billing State/Province (text only)"
TEXT_ONLY_COUNTRY_COL = "Mailing Country (text only)"

CANONICAL_COLUMNS = [
    "Customer Number",
    "Business",
    "First Name",
    "Last Name",
    "Address Line 1",
    "City",
    "State",
    "ZIP Code",
    "Country",
    "Primary Segment",
    "Secondary Segment",
]

SOURCE_HEADERS = [
    "Customer Number",
    "Account Name",
    "First Name",
    "Last Name",
    "Account Source",
    "Account Owner",
    "Contact Owner",
    "Account: Created Date↑",
    "Account: Created Date",
    "Customer Status",
    "Billing Address Line 1",
    "Billing City",
    "Billing State/Province",
    TEXT_ONLY_STATE_COL,
    "Billing Zip/Postal Code",
    "Mailing Country",
    TEXT_ONLY_COUNTRY_COL,
    "Primary Segment",
    "Secondary Segment",
    "Social Profile: Instagram",
]

HEADER_REQUIRED_CORE = {
    "Customer Number",
    "Account Name",
    "First Name",
    "Last Name",
    "Billing Address Line 1",
    "Billing City",
    "Billing Zip/Postal Code",
    "Primary Segment",
    "Secondary Segment",
}

HEADER_STATE_OPTIONS = {"Billing State/Province", TEXT_ONLY_STATE_COL}
HEADER_COUNTRY_OPTIONS = {"Mailing Country", TEXT_ONLY_COUNTRY_COL}


def _normalized_header_key(value):
    text = "" if value is None else str(value)
    text = text.replace("\ufeff", "").strip()
    text = re.sub(r"\s+", " ", text)
    return text.casefold()


SOURCE_HEADER_BY_KEY = {
    _normalized_header_key(header): header
    for header in SOURCE_HEADERS
}


@dataclass(frozen=True)
class ResolvedSource:
    path: Path
    provenance: str
    logical_name: str
    extension: str
    sort_key: tuple


@dataclass(frozen=True)
class TableParseMetadata:
    header_row: int
    preamble_text: str
    removed_blank_columns: tuple
    removed_counter_columns: tuple
    repeated_header_rows: int
    discarded_nondata_rows: int


@dataclass(frozen=True)
class SourceReport:
    logical_name: str
    provenance: str
    source_format: str
    version: str
    sheet_name: str
    header_row: int
    row_count: int
    evidence: tuple
    removed_blank_columns: tuple
    removed_counter_columns: tuple
    repeated_header_rows: int
    discarded_nondata_rows: int

    def detection_payload(self):
        return {
            "name": self.logical_name,
            "format": self.source_format,
            "version": self.version,
            "sheet": self.sheet_name,
            "header_row": self.header_row,
            "rows": self.row_count,
            "evidence": list(self.evidence),
            "removed_blank_columns": list(self.removed_blank_columns),
            "removed_counter_columns": list(self.removed_counter_columns),
            "repeated_header_rows": self.repeated_header_rows,
            "discarded_nondata_rows": self.discarded_nondata_rows,
        }


class WorkbookAdapter:
    def __init__(self, source):
        self.source = source
        self.extension = source.extension
        try:
            if self.extension == ".xlsx":
                self._workbook = openpyxl.load_workbook(
                    source.path,
                    read_only=True,
                    data_only=True,
                )
                self.sheet_names = list(self._workbook.sheetnames)
            elif self.extension == ".xls":
                try:
                    import xlrd
                except ImportError as exc:
                    raise RuntimeError(
                        "Legacy .xls support requires the installed 'xlrd' package."
                    ) from exc
                self._workbook = xlrd.open_workbook(str(source.path), on_demand=True)
                self.sheet_names = list(self._workbook.sheet_names())
            else:
                raise ValueError(f"Unsupported workbook format: {self.extension}")
        except Exception as exc:
            raise ValueError(
                f"Unable to open FOUR HANDS workbook '{source.provenance}' "
                f"({self.extension}): {exc}"
            ) from exc

    def iter_rows(self, sheet_name):
        if self.extension == ".xlsx":
            return self._workbook[sheet_name].iter_rows(values_only=True)

        sheet = self._workbook.sheet_by_name(sheet_name)
        return (sheet.row_values(row_index) for row_index in range(sheet.nrows))

    def close(self):
        if self.extension == ".xlsx":
            self._workbook.close()
        else:
            self._workbook.release_resources()


def current_timestamp():
    return time.strftime("%Y-%m-%dT%H:%M:%S")


def _source_provenance(source):
    if isinstance(source, ResolvedSource):
        return source.provenance
    return str(source)


def _source_logical_name(source):
    if isinstance(source, ResolvedSource):
        return source.logical_name
    return Path(source).name


def parse_job_args():
    job_number = str(sys.argv[1]).strip() if len(sys.argv) > 1 else ""
    drop_number = str(sys.argv[2]).strip() if len(sys.argv) > 2 else "1"
    if not drop_number:
        drop_number = "1"
    return job_number, drop_number


def build_manifest(versions_present, source_reports=None):
    job_number, drop_number = parse_job_args()
    now = current_timestamp()
    manifest = {
        "job_number": job_number,
        "drop_number": drop_number,
        "versions_present": versions_present,
        "versions_complete": [],
        "outputs": {},
        "created_at": now,
        "updated_at": now,
    }
    if source_reports:
        manifest["sources"] = [report.detection_payload() for report in source_reports]
    return manifest


def _stringify_value(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_customer_number(value):
    text = _stringify_value(value)
    if re.fullmatch(r"[0-9]+", text):
        return text
    try:
        number = float(text)
        if number.is_integer():
            return str(int(number))
    except (ValueError, TypeError):
        pass
    return text


def is_valid_customer_number(value):
    try:
        str_value = _stringify_value(value)
        if not str_value:
            return False
        float_value = float(str_value)
        return float_value.is_integer()
    except (ValueError, TypeError):
        return False


def is_valid_value(value):
    value = _stringify_value(value)
    return value.casefold() not in ["", "-", ".", ",", "nan", "none"]


def handle_text_only_column(df, main_col, text_only_col):
    has_main = main_col in df.columns
    has_text_only = text_only_col in df.columns
    if has_text_only:
        if not has_main:
            df = df.rename(columns={text_only_col: main_col})
            df[main_col] = df[main_col].apply(_stringify_value)
            invalid_mask = ~df[main_col].apply(is_valid_value)
            df.loc[invalid_mask, main_col] = ""
        else:
            df[main_col] = df[main_col].apply(_stringify_value)
            df[text_only_col] = df[text_only_col].apply(_stringify_value)
            blank_main_mask = ~df[main_col].apply(is_valid_value)
            valid_text_only_mask = df[text_only_col].apply(is_valid_value)
            rows_to_update = blank_main_mask & valid_text_only_mask
            if rows_to_update.sum() > 0:
                print(f"Transferring {rows_to_update.sum()} values from {text_only_col} to {main_col}")
            df.loc[rows_to_update, main_col] = df.loc[rows_to_update, text_only_col]
        df = df.drop(columns=[text_only_col], errors="ignore")
    return df


def _normalize_sheet_name(value):
    return "".join(ch.lower() for ch in str(value) if ch.isalnum())


def _is_ignore_marker(sheet_name):
    return IGNORE_MARKER_RE.fullmatch(str(sheet_name).strip()) is not None


def eligible_sheet_names(sheet_names):
    eligible = []
    for sheet_name in sheet_names:
        if _is_ignore_marker(sheet_name):
            break
        eligible.append(sheet_name)
    return eligible


def resolve_sheet_name(sheet_names, desired_name):
    """Resolve a desired name only within the already eligible sheet set."""
    if desired_name in sheet_names:
        return desired_name

    desired_lower = desired_name.lower()
    for sheet_name in sheet_names:
        if sheet_name.lower() == desired_lower:
            return sheet_name

    desired_norm = _normalize_sheet_name(desired_name)
    for sheet_name in sheet_names:
        if _normalize_sheet_name(sheet_name) == desired_norm:
            return sheet_name

    tokens = [
        token
        for token in re.split(r"[^A-Za-z0-9]+", desired_name.lower())
        if token
    ]
    for sheet_name in sheet_names:
        sheet_lower = sheet_name.lower()
        if all(token in sheet_lower for token in tokens):
            return sheet_name

    return None


def _matches_standalone_hospitality_filename(source):
    stem = Path(_source_logical_name(source)).stem
    normalized_stem = re.sub(r"[\s_-]+", " ", stem).strip()
    return STANDALONE_HOSPITALITY_RE.match(normalized_stem) is not None


def _cell_has_value(value):
    return value is not None and str(value).strip() != ""


def _clean_header_value(value):
    text = _stringify_value(value).replace("\ufeff", "")
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return ""
    return SOURCE_HEADER_BY_KEY.get(_normalized_header_key(text), text)


def _looks_like_header(row):
    headers = {
        _clean_header_value(value)
        for value in row
        if _clean_header_value(value)
    }
    return (
        HEADER_REQUIRED_CORE.issubset(headers)
        and bool(headers & HEADER_STATE_OPTIONS)
        and bool(headers & HEADER_COUNTRY_OPTIONS)
    )


def _excel_column_name(index):
    number = index + 1
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def _is_sequential_counter(values):
    if not values or any(not _cell_has_value(value) for value in values):
        return False
    integers = []
    for value in values:
        try:
            number = float(str(value).strip())
        except (ValueError, TypeError):
            return False
        if not number.is_integer():
            return False
        integers.append(int(number))
    if len(integers) == 1:
        return integers[0] in (0, 1)
    return all(current == previous + 1 for previous, current in zip(integers, integers[1:]))


def _row_matches_header(row, header):
    for index, header_value in enumerate(header):
        if not header_value:
            continue
        value = row[index] if index < len(row) else None
        if _clean_header_value(value) != header_value:
            return False
    return True


def _validate_and_order_schema(df, source, sheet_name):
    provenance = _source_provenance(source)
    duplicate_columns = [
        str(column)
        for column in df.columns[df.columns.duplicated()].tolist()
    ]
    if duplicate_columns:
        raise ValueError(
            f"Schema mismatch in '{provenance}' / '{sheet_name}': "
            f"duplicate columns after normalization: {duplicate_columns}"
        )

    missing = [column for column in CANONICAL_COLUMNS if column not in df.columns]
    unexpected = [column for column in df.columns if column not in CANONICAL_COLUMNS]
    if missing or unexpected:
        details = []
        if missing:
            details.append(f"missing required columns: {missing}")
        if unexpected:
            details.append(f"unexpected columns: {unexpected}")
        raise ValueError(
            f"Schema mismatch in '{provenance}' / '{sheet_name}': "
            + "; ".join(details)
        )

    return df.loc[:, CANONICAL_COLUMNS].copy()


def read_and_normalize_sheet(
    source,
    sheet_name,
    rows,
    allow_missing_header=False,
):
    provenance = _source_provenance(source)
    header = None
    header_row = None
    candidate_rows = []
    preamble_values = []
    trailing_blank_rows = 0

    for row_number, raw_row in enumerate(rows, start=1):
        row = list(raw_row)
        row_has_content = any(_cell_has_value(value) for value in row)
        if header is None:
            if row_number > MAX_HEADER_SCAN_ROWS:
                break
            if _looks_like_header(row):
                header = [_clean_header_value(value) for value in row]
                header_row = row_number
            elif row_has_content:
                preamble_values.extend(
                    _stringify_value(value)
                    for value in row
                    if _cell_has_value(value)
                )
            continue

        if not row_has_content:
            if candidate_rows:
                trailing_blank_rows += 1
                if trailing_blank_rows >= MAX_TRAILING_BLANK_ROWS:
                    break
            continue

        trailing_blank_rows = 0
        if len(row) > len(header) and any(
            _cell_has_value(value) for value in row[len(header) :]
        ):
            raise ValueError(
                f"Schema mismatch in '{provenance}' / '{sheet_name}' at row {row_number}: "
                "data exists beyond the detected header width."
            )
        normalized_row = row[: len(header)]
        if len(normalized_row) < len(header):
            normalized_row.extend([None] * (len(header) - len(normalized_row)))
        candidate_rows.append((row_number, normalized_row))

    if header is None:
        if allow_missing_header:
            return None, None
        raise ValueError(
            f"A complete FOUR HANDS header was not found within the first "
            f"{MAX_HEADER_SCAN_ROWS} rows of '{provenance}' / '{sheet_name}'."
        )

    duplicate_source_columns = [
        column for column in set(header) if column and header.count(column) > 1
    ]
    if duplicate_source_columns:
        raise ValueError(
            f"Schema mismatch in '{provenance}' / '{sheet_name}': "
            f"duplicate source columns: {sorted(duplicate_source_columns)}"
        )

    customer_index = header.index("Customer Number")
    valid_rows = []
    repeated_header_rows = 0
    discarded_nondata_rows = 0
    for row_number, row in candidate_rows:
        if _row_matches_header(row, header):
            repeated_header_rows += 1
            continue
        if _clean_header_value(row[customer_index]) == "Customer Number":
            raise ValueError(
                f"Partial or altered repeated header found in '{provenance}' / "
                f"'{sheet_name}' at row {row_number}."
            )
        if is_valid_customer_number(row[customer_index]):
            valid_rows.append((row_number, row))
        else:
            discarded_nondata_rows += 1

    removed_blank_columns = []
    removed_counter_columns = []
    unnamed_indices = [index for index, value in enumerate(header) if not value]
    for index in unnamed_indices:
        values = [row[index] for _, row in valid_rows]
        column_name = _excel_column_name(index)
        if not any(_cell_has_value(value) for value in values):
            removed_blank_columns.append(column_name)
            continue
        if _is_sequential_counter(values):
            removed_counter_columns.append(column_name)
            continue
        sample_rows = [
            str(row_number)
            for row_number, row in valid_rows
            if _cell_has_value(row[index])
        ][:5]
        raise ValueError(
            f"Unlabeled column {column_name} in '{provenance}' / '{sheet_name}' "
            f"contains customer data at row(s) {', '.join(sample_rows)}. "
            "The column was not removed because its meaning is unknown."
        )

    kept_indices = [index for index, value in enumerate(header) if value]
    data_rows = [
        [row[index] for index in kept_indices]
        for _, row in valid_rows
    ]
    kept_header = [header[index] for index in kept_indices]
    df = pd.DataFrame(data_rows, columns=kept_header)
    if "Customer Number" in df.columns:
        df["Customer Number"] = df["Customer Number"].apply(_normalize_customer_number)
    df = df.rename(columns=COLUMN_RENAME)

    df = handle_text_only_column(df, "State", TEXT_ONLY_STATE_COL)
    df = handle_text_only_column(df, "Country", TEXT_ONLY_COUNTRY_COL)
    df = df.drop(
        columns=[column for column in COLUMNS_TO_REMOVE if column in df.columns],
        errors="ignore",
    )

    for column in df.columns:
        df[column] = df[column].apply(_stringify_value).apply(
            lambda value: value.encode("utf-8", "ignore").decode("utf-8", "ignore")
        )

    df = df.dropna(how="all").reset_index(drop=True)
    metadata = TableParseMetadata(
        header_row=header_row,
        preamble_text=" ".join(preamble_values),
        removed_blank_columns=tuple(removed_blank_columns),
        removed_counter_columns=tuple(removed_counter_columns),
        repeated_header_rows=repeated_header_rows,
        discarded_nondata_rows=discarded_nondata_rows,
    )
    return _validate_and_order_schema(df, source, sheet_name), metadata


def _signals_from_text(text, label):
    signals = set()
    evidence = []
    text = _stringify_value(text)
    if not text:
        return signals, evidence
    if RESIDENTIAL_SIGNAL_RE.search(text):
        signals.add(VERSION_RESIDENTIAL)
        evidence.append(f"{label} indicates Residential/Brand")
    if HOSPITALITY_SIGNAL_RE.search(text):
        signals.add(VERSION_HOSPITALITY)
        evidence.append(f"{label} indicates Hospitality/Commercial")
    return signals, evidence


def _segment_signals(df):
    signals = set()
    evidence = []
    for column in ("Primary Segment", "Secondary Segment"):
        if column not in df.columns:
            continue
        values = sorted(
            {
                _stringify_value(value)
                for value in df[column].tolist()
                if _stringify_value(value)
            }
        )
        for value in values:
            value_signals, _ = _signals_from_text(value, column)
            for version in value_signals:
                signals.add(version)
                evidence.append(f"{column} value '{value}'")
    return signals, evidence


def _single_signal_or_error(signals, provenance, description):
    if len(signals) > 1:
        raise ValueError(
            f"Conflicting {description} signals in FOUR HANDS source "
            f"'{provenance}': {', '.join(sorted(signals))}."
        )
    return next(iter(signals), "")


def _resolve_table_version(source, sheet_name, df, metadata, expected_version=""):
    provenance = _source_provenance(source)
    segment_signals, segment_evidence = _segment_signals(df)
    segment_version = _single_signal_or_error(
        segment_signals,
        provenance,
        "Primary/Secondary Segment",
    )

    filename_signals, filename_evidence = _signals_from_text(
        _source_logical_name(source),
        "filename",
    )
    sheet_signals, sheet_evidence = _signals_from_text(
        sheet_name,
        "worksheet name",
    )
    name_signals = filename_signals | sheet_signals
    name_evidence = filename_evidence + sheet_evidence

    if expected_version:
        sheet_version = _single_signal_or_error(
            sheet_signals,
            provenance,
            "worksheet name",
        )
        conflicting = {
            version
            for version in (segment_version, sheet_version)
            if version and version != expected_version
        }
        if conflicting:
            raise ValueError(
                f"FOUR HANDS version conflict in '{provenance}' / '{sheet_name}': "
                f"legacy worksheet classification is {expected_version}, but content or "
                f"naming indicates {', '.join(sorted(conflicting))}."
            )
        evidence = [f"legacy worksheet '{sheet_name}'"]
        evidence.extend(segment_evidence)
        evidence.extend(sheet_evidence)
        if expected_version in filename_signals:
            evidence.extend(filename_evidence)
        return expected_version, tuple(dict.fromkeys(evidence))

    name_version = _single_signal_or_error(name_signals, provenance, "filename/worksheet")

    if segment_version and name_version and segment_version != name_version:
        raise ValueError(
            f"FOUR HANDS version conflict in '{provenance}' / '{sheet_name}': "
            f"segment data indicates {segment_version}, but filename/worksheet naming "
            f"indicates {name_version}."
        )

    resolved_version = segment_version or name_version
    evidence = segment_evidence + name_evidence
    if not resolved_version:
        preamble_signals, preamble_evidence = _signals_from_text(
            metadata.preamble_text,
            "report preamble",
        )
        resolved_version = _single_signal_or_error(
            preamble_signals,
            provenance,
            "report preamble",
        )
        evidence.extend(preamble_evidence)

    if not resolved_version:
        raise ValueError(
            f"Unable to classify FOUR HANDS source '{provenance}' / '{sheet_name}' "
            "as RESIDENTIAL or HOSPITALITY. No unambiguous version signal was found "
            "in the worksheet name, filename, report preamble, or segment data."
        )
    return resolved_version, tuple(dict.fromkeys(evidence))


def _workbook_format_label(source, metadata, legacy=False):
    if legacy:
        return "LEGACY_WORKBOOK"
    if _matches_standalone_hospitality_filename(source) and metadata.header_row == 1:
        return "LEGACY_STANDALONE"
    if (
        metadata.header_row > 1
        or metadata.removed_blank_columns
        or metadata.removed_counter_columns
    ):
        return "LOOK_BOOK_EXPORT"
    return "GENERIC_WORKBOOK"


def _build_source_report(source, source_format, version, sheet_name, df, metadata, evidence):
    return SourceReport(
        logical_name=_source_logical_name(source),
        provenance=_source_provenance(source),
        source_format=source_format,
        version=version,
        sheet_name=sheet_name,
        header_row=metadata.header_row,
        row_count=len(df),
        evidence=evidence,
        removed_blank_columns=metadata.removed_blank_columns,
        removed_counter_columns=metadata.removed_counter_columns,
        repeated_header_rows=metadata.repeated_header_rows,
        discarded_nondata_rows=metadata.discarded_nondata_rows,
    )


def _log_source_report(report, verbose):
    if not verbose:
        return
    cleanup = []
    if report.removed_blank_columns:
        cleanup.append("blank columns " + ", ".join(report.removed_blank_columns))
    if report.removed_counter_columns:
        cleanup.append("counter columns " + ", ".join(report.removed_counter_columns))
    if report.repeated_header_rows:
        cleanup.append(f"{report.repeated_header_rows} repeated header row(s)")
    if report.discarded_nondata_rows:
        cleanup.append(f"{report.discarded_nondata_rows} non-customer row(s)")
    cleanup_text = f"; removed {', '.join(cleanup)}" if cleanup else ""
    print(
        f"Classified '{report.provenance}' / '{report.sheet_name}' as "
        f"{report.version} ({report.source_format}, header row {report.header_row}): "
        f"{report.row_count} valid rows{cleanup_text}"
    )


def classify_workbook(source, verbose=True):
    residential_dfs = []
    hospitality_dfs = []
    reports = []
    provenance = _source_provenance(source)

    if verbose:
        print(f"Reading: {provenance}")

    workbook = WorkbookAdapter(source)
    try:
        all_sheet_names = workbook.sheet_names
        eligible_names = eligible_sheet_names(all_sheet_names)
        ignored_names = all_sheet_names[len(eligible_names) :]
        if ignored_names and _is_ignore_marker(ignored_names[0]) and verbose:
            print(
                f"Ignoring marker and sheets to its right in '{provenance}': "
                + ", ".join(ignored_names)
            )

        desired_mappings = [
            (SHEET_RESIDENTIAL, VERSION_RESIDENTIAL),
            (SHEET_COMMERCIAL, VERSION_HOSPITALITY),
            (SHEET_NEW_ADDRESSES, VERSION_HOSPITALITY),
        ]
        classified_sheets = {}
        for desired_name, version in desired_mappings:
            actual_name = resolve_sheet_name(eligible_names, desired_name)
            if not actual_name:
                continue
            existing_version = classified_sheets.get(actual_name)
            if existing_version and existing_version != version:
                raise ValueError(
                    f"Conflicting sheet classification in '{provenance}': "
                    f"'{actual_name}' matched multiple versions"
                )
            classified_sheets[actual_name] = version

        if classified_sheets:
            for sheet_name in eligible_names:
                version = classified_sheets.get(sheet_name)
                if not version:
                    continue
                df, metadata = read_and_normalize_sheet(
                    source,
                    sheet_name,
                    workbook.iter_rows(sheet_name),
                )
                if df.empty:
                    if verbose:
                        print(
                            f"WARNING: Eligible FOUR HANDS sheet '{provenance}' / "
                            f"'{sheet_name}' contains no valid Customer Number rows."
                        )
                    continue
                version, evidence = _resolve_table_version(
                    source,
                    sheet_name,
                    df,
                    metadata,
                    expected_version=version,
                )
                if version == VERSION_RESIDENTIAL:
                    residential_dfs.append(df)
                else:
                    hospitality_dfs.append(df)
                report = _build_source_report(
                    source,
                    _workbook_format_label(source, metadata, legacy=True),
                    version,
                    sheet_name,
                    df,
                    metadata,
                    evidence,
                )
                reports.append(report)
                _log_source_report(report, verbose)
            if not reports:
                raise ValueError(
                    f"No valid customer rows were found in recognized FOUR HANDS "
                    f"worksheets in '{provenance}'."
                )
            return residential_dfs, hospitality_dfs, reports

        if not eligible_names:
            raise ValueError(
                f"FOUR HANDS workbook '{provenance}' has no eligible worksheets "
                "before the IGNORE >>> marker."
            )

        table_candidates = []
        for sheet_name in eligible_names:
            df, metadata = read_and_normalize_sheet(
                source,
                sheet_name,
                workbook.iter_rows(sheet_name),
                allow_missing_header=True,
            )
            if df is None:
                continue
            if df.empty:
                if verbose:
                    print(
                        f"WARNING: FOUR HANDS table '{provenance}' / '{sheet_name}' "
                        "contains no valid Customer Number rows."
                    )
                continue
            table_candidates.append((sheet_name, df, metadata))

        if not table_candidates:
            raise ValueError(
                f"No complete FOUR HANDS customer table was found in '{provenance}'."
            )

        for sheet_name, df, metadata in table_candidates:
            version, evidence = _resolve_table_version(
                source,
                sheet_name,
                df,
                metadata,
            )
            if version == VERSION_RESIDENTIAL:
                residential_dfs.append(df)
            else:
                hospitality_dfs.append(df)
            report = _build_source_report(
                source,
                _workbook_format_label(source, metadata),
                version,
                sheet_name,
                df,
                metadata,
                evidence,
            )
            reports.append(report)
            _log_source_report(report, verbose)
        return residential_dfs, hospitality_dfs, reports
    finally:
        workbook.close()


def classify_csv_source(source, verbose=True):
    provenance = _source_provenance(source)
    try:
        with source.path.open("r", encoding="utf-8-sig", newline="") as csv_file:
            df, metadata = read_and_normalize_sheet(
                source,
                "<CSV>",
                csv.reader(csv_file),
            )
    except UnicodeDecodeError as exc:
        raise ValueError(
            f"Unable to decode FOUR HANDS CSV source '{provenance}' as UTF-8: {exc}"
        ) from exc
    except OSError as exc:
        raise ValueError(
            f"Unable to read FOUR HANDS CSV source '{provenance}': {exc}"
        ) from exc

    if df.empty:
        raise ValueError(
            f"FOUR HANDS CSV source '{provenance}' contains no valid Customer "
            "Number rows."
        )

    version, evidence = _resolve_table_version(
        source,
        "<CSV>",
        df,
        metadata,
    )
    report = _build_source_report(
        source,
        "CSV",
        version,
        "<CSV>",
        df,
        metadata,
        evidence,
    )
    _log_source_report(report, verbose)
    if version == VERSION_RESIDENTIAL:
        return [df], [], [report]
    return [], [df], [report]


def _record_resolution_warning(message, warnings, verbose):
    warnings.append(message)
    if verbose:
        print(f"WARNING: {message}")


def list_raw_source_artifacts(source_dir):
    source_dir = Path(source_dir)
    if not source_dir.exists() or not source_dir.is_dir():
        return []
    return sorted(
        [
            path
            for path in source_dir.iterdir()
            if path.is_file() and path.suffix.casefold() in SUPPORTED_RAW_EXTENSIONS
        ],
        key=lambda path: (path.name.casefold(), str(path).casefold()),
    )


def _safe_zip_member_parts(member_name, provenance):
    if "\x00" in member_name:
        raise ValueError(f"Unsafe ZIP member contains a NUL character: {provenance}")

    normalized = member_name.replace("\\", "/")
    if normalized.startswith("/") or normalized.startswith("//"):
        raise ValueError(f"Unsafe absolute ZIP member path: {provenance}")
    if re.match(r"^[A-Za-z]:", normalized):
        raise ValueError(f"Unsafe drive-qualified ZIP member path: {provenance}")

    parts = []
    for part in normalized.split("/"):
        if part in ("", "."):
            continue
        if part == "..":
            raise ValueError(f"Unsafe path-traversal ZIP member: {provenance}")
        if ":" in part:
            raise ValueError(f"Unsafe ZIP member path component: {provenance}")
        if part.endswith((" ", ".")):
            raise ValueError(
                f"Unsafe ZIP member path component with trailing space/dot: {provenance}"
            )
        parts.append(part)

    if not parts:
        return []
    return parts


def _resolve_zip_payloads(raw_zip, raw_index, temp_root, warnings, verbose):
    archive_root = temp_root / f"archive_{raw_index:04d}"
    archive_root.mkdir(parents=True, exist_ok=False)
    resolved = []
    seen_destinations = set()
    nested_zip_provenances = []
    unsupported_provenances = []

    try:
        with zipfile.ZipFile(raw_zip, "r") as archive:
            members = sorted(
                archive.infolist(),
                key=lambda info: info.filename.replace("\\", "/").casefold(),
            )
            for member in members:
                member_path = member.filename.replace("\\", "/")
                provenance = f"{raw_zip}!/{member_path}"
                parts = _safe_zip_member_parts(member.filename, provenance)
                if member.is_dir() or not parts:
                    continue

                destination_key = "/".join(parts).casefold()
                if destination_key in seen_destinations:
                    raise ValueError(
                        f"Case-insensitive ZIP destination collision in '{raw_zip}': "
                        f"'{member_path}'"
                    )
                seen_destinations.add(destination_key)

                extension = Path(parts[-1]).suffix.casefold()
                if extension == ".zip":
                    nested_zip_provenances.append(provenance)
                    continue
                if extension not in SUPPORTED_SOURCE_EXTENSIONS:
                    unsupported_provenances.append(provenance)
                    continue

                destination = archive_root.joinpath(*parts)
                resolved_archive_root = archive_root.resolve()
                resolved_destination = destination.resolve(strict=False)
                try:
                    resolved_destination.relative_to(resolved_archive_root)
                except ValueError as exc:
                    raise ValueError(
                        f"Unsafe ZIP member resolves outside temporary extraction root: {provenance}"
                    ) from exc

                destination.parent.mkdir(parents=True, exist_ok=True)
                try:
                    with archive.open(member, "r") as source_file:
                        with destination.open("xb") as destination_file:
                            shutil.copyfileobj(source_file, destination_file)
                except Exception as exc:
                    raise ValueError(
                        f"Unable to extract FOUR HANDS ZIP payload '{provenance}': {exc}"
                    ) from exc

                resolved.append(
                    ResolvedSource(
                        path=destination,
                        provenance=provenance,
                        logical_name=parts[-1],
                        extension=extension,
                        sort_key=(raw_zip.name.casefold(), member_path.casefold()),
                    )
                )
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Unable to read FOUR HANDS ZIP archive '{raw_zip}': {exc}") from exc
    except OSError as exc:
        raise ValueError(f"Unable to resolve FOUR HANDS ZIP archive '{raw_zip}': {exc}") from exc

    for provenance in nested_zip_provenances[:MAX_REPORTED_ZIP_WARNINGS]:
        _record_resolution_warning(
            f"Nested ZIP payload ignored (recursive extraction is not supported): {provenance}",
            warnings,
            verbose,
        )
    if len(nested_zip_provenances) > MAX_REPORTED_ZIP_WARNINGS:
        _record_resolution_warning(
            f"{len(nested_zip_provenances) - MAX_REPORTED_ZIP_WARNINGS} additional nested ZIP payload(s) were ignored in {raw_zip}.",
            warnings,
            verbose,
        )

    for provenance in unsupported_provenances[:MAX_REPORTED_ZIP_WARNINGS]:
        _record_resolution_warning(
            f"Unsupported ZIP payload ignored: {provenance}",
            warnings,
            verbose,
        )
    if len(unsupported_provenances) > MAX_REPORTED_ZIP_WARNINGS:
        _record_resolution_warning(
            f"{len(unsupported_provenances) - MAX_REPORTED_ZIP_WARNINGS} additional unsupported ZIP payload(s) were ignored in {raw_zip}.",
            warnings,
            verbose,
        )
    if not resolved:
        _record_resolution_warning(
            f"ZIP archive contains no supported FOUR HANDS payloads: {raw_zip}",
            warnings,
            verbose,
        )
    return resolved


@contextmanager
def resolve_source_artifacts(source_dir, verbose=True, warnings=None):
    warnings = warnings if warnings is not None else []
    raw_artifacts = list_raw_source_artifacts(source_dir)
    with tempfile.TemporaryDirectory(prefix="goji_fourhands_sources_") as temp_dir:
        temp_root = Path(temp_dir)
        resolved_sources = []
        for raw_index, raw_artifact in enumerate(raw_artifacts):
            extension = raw_artifact.suffix.casefold()
            if extension == ".zip":
                resolved_sources.extend(
                    _resolve_zip_payloads(
                        raw_artifact,
                        raw_index,
                        temp_root,
                        warnings,
                        verbose,
                    )
                )
                continue

            resolved_sources.append(
                ResolvedSource(
                    path=raw_artifact,
                    provenance=str(raw_artifact),
                    logical_name=raw_artifact.name,
                    extension=extension,
                    sort_key=(raw_artifact.name.casefold(), ""),
                )
            )

        resolved_sources.sort(key=lambda source: source.sort_key)
        yield raw_artifacts, resolved_sources


def _sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_unique_source_hashes(sources):
    paths_by_hash = {}
    for source in sources:
        try:
            digest = _sha256_file(source.path)
        except OSError as exc:
            raise ValueError(
                f"Unable to hash FOUR HANDS source '{source.provenance}': {exc}"
            ) from exc
        paths_by_hash.setdefault(digest, []).append(source.provenance)

    duplicate_groups = [
        (digest, paths)
        for digest, paths in paths_by_hash.items()
        if len(paths) > 1
    ]
    if not duplicate_groups:
        return

    details = []
    for digest, paths in duplicate_groups:
        details.append(
            f"SHA-256 {digest}: " + ", ".join(paths)
        )
    raise ValueError(
        "Exact duplicate FOUR HANDS resolved customer-data sources were found, possibly "
        "across direct files or ZIP payloads. Remove the duplicate before retrying. "
        + " | ".join(details)
    )


def classify_sources(sources, verbose=True):
    residential_dfs = []
    hospitality_dfs = []
    reports = []
    for source in sorted(sources, key=lambda item: item.sort_key):
        if source.extension == ".csv":
            file_residential, file_hospitality, file_reports = classify_csv_source(
                source,
                verbose=verbose,
            )
        else:
            file_residential, file_hospitality, file_reports = classify_workbook(
                source,
                verbose=verbose,
            )
        residential_dfs.extend(file_residential)
        hospitality_dfs.extend(file_hospitality)
        reports.extend(file_reports)
    return residential_dfs, hospitality_dfs, reports


def classify_source_directory(source_dir, verbose=True, warnings=None):
    with resolve_source_artifacts(
        source_dir,
        verbose=verbose,
        warnings=warnings,
    ) as (raw_artifacts, resolved_sources):
        if not raw_artifacts:
            raise ValueError(f"No supported FOUR HANDS source artifacts found: {source_dir}")
        if not resolved_sources:
            raise ValueError(
                f"No supported FOUR HANDS customer-data sources were resolved from: {source_dir}"
            )
        validate_unique_source_hashes(resolved_sources)
        residential_dfs, hospitality_dfs, reports = classify_sources(
            resolved_sources,
            verbose=verbose,
        )
        provenances = [source.provenance for source in resolved_sources]
    return provenances, residential_dfs, hospitality_dfs, reports


def detected_versions(residential_dfs, hospitality_dfs):
    versions = []
    if residential_dfs:
        versions.append(VERSION_RESIDENTIAL)
    if hospitality_dfs:
        versions.append(VERSION_HOSPITALITY)
    return versions


def emit_detection_result(status, versions=None, error="", warnings=None, source_reports=None):
    result = {
        "status": status,
        "versions": versions or [],
    }
    if error:
        result["error"] = error
    if warnings:
        result["warnings"] = warnings
    if source_reports:
        result["sources"] = [
            report.detection_payload()
            for report in source_reports
        ]

    print(DETECTION_BEGIN_MARKER)
    print(json.dumps(result, separators=(",", ":")))
    print(DETECTION_END_MARKER)


def detection_source_from_args(args):
    if args == [DETECTION_MODE]:
        return SOURCE_DIR
    if (
        len(args) == 3
        and args[0] == DETECTION_MODE
        and args[1] == DETECTION_SOURCE_OPTION
    ):
        return Path(args[2])
    raise ValueError(
        f"Usage: {os.path.basename(sys.argv[0])} {DETECTION_MODE} "
        f"[{DETECTION_SOURCE_OPTION} <directory>]"
    )


def run_detection_only(args):
    warnings = []
    try:
        source_dir = detection_source_from_args(args)
        _, residential_dfs, hospitality_dfs, source_reports = classify_source_directory(
            source_dir,
            verbose=False,
            warnings=warnings,
        )
        versions = detected_versions(residential_dfs, hospitality_dfs)
        if not versions:
            raise ValueError(
                "No supported non-empty FOUR HANDS version data was detected."
            )
        emit_detection_result(
            "ok",
            versions,
            warnings=warnings,
            source_reports=source_reports,
        )
        return 0
    except Exception as exc:
        emit_detection_result("error", error=str(exc), warnings=warnings)
        traceback.print_exc()
        return 1


def _normalized_customer_key(value):
    try:
        number = float(str(value).strip())
        if number.is_integer():
            return str(int(number))
    except (ValueError, TypeError):
        pass
    return str(value).strip()


def _warn_customer_number_duplicates(df, version):
    keys = df["Customer Number"].apply(_normalized_customer_key)
    duplicate_mask = keys.duplicated(keep=False)
    if not duplicate_mask.any():
        return
    duplicate_keys = keys[duplicate_mask]
    print(
        f"WARNING: {version} contains {len(duplicate_keys)} rows across "
        f"{duplicate_keys.nunique()} duplicate Customer Number value(s). "
        "Duplicate records are preserved."
    )


def build_combined_outputs(residential_dfs, hospitality_dfs):
    buckets = {
        VERSION_RESIDENTIAL: residential_dfs,
        VERSION_HOSPITALITY: hospitality_dfs,
    }
    combined_outputs = {}
    for version in VERSION_ORDER:
        dataframes = buckets[version]
        if not dataframes:
            continue
        combined = pd.concat(dataframes, ignore_index=True).dropna(how="all")
        combined = combined.loc[:, CANONICAL_COLUMNS].reset_index(drop=True)
        if combined.empty:
            continue
        _warn_customer_number_duplicates(combined, version)
        combined_outputs[version] = combined
    return combined_outputs


def _stage_dataframe(df, final_path):
    final_path = Path(final_path)
    final_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temp_name = tempfile.mkstemp(
        prefix=".goji_fourhands_",
        suffix=".csv.tmp",
        dir=str(final_path.parent),
    )
    temp_path = Path(temp_name)
    try:
        with os.fdopen(
            file_descriptor,
            "w",
            encoding="utf-8-sig",
            newline="",
        ) as staged_file:
            df.to_csv(staged_file, index=False)
            staged_file.flush()
            os.fsync(staged_file.fileno())

        with temp_path.open("r", encoding="utf-8-sig", newline="") as check_file:
            reader = csv.reader(check_file)
            header = next(reader, None)
            rows = list(reader)
        if header != CANONICAL_COLUMNS:
            raise ValueError(
                f"Staged output header verification failed for {final_path}: {header}"
            )
        if len(rows) != len(df):
            raise ValueError(
                f"Staged output row-count verification failed for {final_path}: "
                f"expected {len(df)}, found {len(rows)}"
            )
        if any(row == header for row in rows):
            raise ValueError(
                f"Staged output contains an embedded duplicate header row: {final_path}"
            )
        if temp_path.stat().st_size <= 0:
            raise ValueError(f"Staged output is empty: {final_path}")
        return temp_path
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def _stage_manifest(manifest, manifest_path):
    manifest_path = Path(manifest_path)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temp_name = tempfile.mkstemp(
        prefix=".goji_fourhands_",
        suffix=".json.tmp",
        dir=str(manifest_path.parent),
    )
    temp_path = Path(temp_name)
    try:
        with os.fdopen(file_descriptor, "w", encoding="utf-8", newline="") as staged_file:
            json.dump(manifest, staged_file, indent=2)
            staged_file.write("\n")
            staged_file.flush()
            os.fsync(staged_file.fileno())
        loaded = json.loads(temp_path.read_text(encoding="utf-8"))
        if loaded != manifest:
            raise ValueError("Staged FOUR HANDS manifest verification failed")
        return temp_path
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def commit_generated_state(
    combined_outputs,
    manifest,
    res_output_file=RES_OUTPUT_FILE,
    hosp_output_file=HOSP_OUTPUT_FILE,
    manifest_file=MANIFEST_FILE,
):
    output_paths = {
        VERSION_RESIDENTIAL: Path(res_output_file),
        VERSION_HOSPITALITY: Path(hosp_output_file),
    }
    staged_outputs = {}
    staged_manifest = None
    try:
        for version in VERSION_ORDER:
            dataframe = combined_outputs.get(version)
            if dataframe is not None:
                staged_outputs[version] = _stage_dataframe(
                    dataframe,
                    output_paths[version],
                )
        staged_manifest = _stage_manifest(manifest, manifest_file)

        for version in VERSION_ORDER:
            staged_path = staged_outputs.get(version)
            if staged_path is not None:
                os.replace(staged_path, output_paths[version])
                print(f"{version} OUTPUT: {output_paths[version]}")

        for version in VERSION_ORDER:
            if version not in combined_outputs and output_paths[version].exists():
                output_paths[version].unlink()
                print(f"Removed stale generated {version} input: {output_paths[version]}")

        os.replace(staged_manifest, manifest_file)
        print(f"Manifest created: {manifest_file}")
    finally:
        for staged_path in staged_outputs.values():
            Path(staged_path).unlink(missing_ok=True)
        if staged_manifest is not None:
            Path(staged_manifest).unlink(missing_ok=True)


def process_initial(
    source_dir=SOURCE_DIR,
    res_output_file=RES_OUTPUT_FILE,
    hosp_output_file=HOSP_OUTPUT_FILE,
    manifest_file=MANIFEST_FILE,
):
    print("=== READING FILES ===")
    resolved_sources, residential_dfs, hospitality_dfs, source_reports = classify_source_directory(
        source_dir,
        verbose=True,
    )
    if not resolved_sources:
        raise ValueError(f"No resolved FOUR HANDS sources found in source folder: {source_dir}")

    print("=== MERGING FILES ===")
    combined_outputs = build_combined_outputs(
        residential_dfs,
        hospitality_dfs,
    )
    versions_present = [
        version for version in VERSION_ORDER if version in combined_outputs
    ]
    if not versions_present:
        raise ValueError(
            "No supported non-empty FOUR HANDS version data was generated. "
            "Expected eligible Residential, Commercial, New Addresses, or an "
            "approved standalone Hospitality workbook."
        )

    manifest = build_manifest(versions_present, source_reports=source_reports)
    print("=== STAGING OUTPUTS ===")
    commit_generated_state(
        combined_outputs,
        manifest,
        res_output_file=res_output_file,
        hosp_output_file=hosp_output_file,
        manifest_file=manifest_file,
    )
    return versions_present, combined_outputs


def main():
    try:
        versions_present, _ = process_initial()
        print("=== COMPLETE ===")
        print("=== NAS_FOLDER_PATH ===")
        print(str(BASE_DIR))
        print("=== END_NAS_FOLDER_PATH ===")
        print(f"VERSIONS GENERATED: {', '.join(versions_present)}")
        print("PROCESS COMPLETE")
        return 0
    except Exception as exc:
        print("FATAL ERROR:", str(exc))
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == DETECTION_MODE:
        sys.exit(run_detection_only(sys.argv[1:]))
    sys.exit(main())
